from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


ai_check = load_module("ai_check", ROOT / "ai-check" / "scripts" / "ai_check.py")
ai_rewrite = load_module("ai_rewrite_docx", ROOT / "ai-check" / "scripts" / "ai_rewrite_docx.py")


def test_trace_features_are_evidence_not_probability():
    features = ai_check.scan_trace_features(
        "首先，本文提出一种高效可靠的系统框架。其次，实验结果表明准确率提升了 3.2% [4]。"
        "最后，综上所述，该方法具有重要意义。"
    )
    data = ai_check.asdict(features)

    assert "riskSignals" in data
    assert "mitigatingSignals" in data
    assert "aiScore" not in data
    assert "score" not in data
    assert any(signal.startswith(("template phrase", "generic claims")) for signal in data["riskSignals"])
    assert any(signal.startswith(("concrete evidence", "quantified results")) for signal in data["mitigatingSignals"])


def test_build_chunks_skips_keywords_titles_and_references():
    chunks = ai_check.build_chunks(
        "\n".join(
            [
                "论文题目",
                "摘要",
                "首先，本文围绕水声目标识别任务展开研究。此外，实验结果表明该方法在公开数据集上准确率提升了 2.1% [4]。",
                "关键词：水声目标识别；深度学习",
                "1 绪论",
                "传统水声目标识别技术主要依赖预先设计的特征描述符和浅层分类器[4]，面临特征设计依赖专家知识等挑战。",
                "参考文献",
                "[4] Some paper.",
            ]
        )
    )

    assert chunks
    assert {chunk.sectionType for chunk in chunks} <= {"abstract", "body", "conclusion"}
    assert all("关键词" not in chunk.text for chunk in chunks)
    assert all("参考文献" not in chunk.text for chunk in chunks)


def test_summarize_title_skips_non_title_front_matter():
    text = "\n".join(
        [
            "摘 要",
            "本文围绕水声目标识别任务展开研究，实验结果表明该方法在公开数据集上准确率提升了 2.1%。",
            "关键词：水声目标识别，深度学习，轻量级网络",
            "ABSTRACT",
            "With the advancement of maritime strategy, UATR has become a research hotspot.",
        ]
    )

    assert ai_check.summarize_title(Path("AI测试.docx"), text) == "AI测试"


def test_summarize_title_uses_probable_cover_title():
    text = "\n".join(
        [
            "基于深度神经网络的轻量化与可解释水声目标识别研究",
            "摘 要",
            "本文围绕水声目标识别任务展开研究，实验结果表明该方法在公开数据集上准确率提升了 2.1%。",
        ]
    )

    assert ai_check.summarize_title(Path("AI测试.docx"), text) == "基于深度神经网络的轻量化与可解释水声目标识别研究"


def test_report_shows_ai_score_and_reason_in_english():
    scan = {
        "documentTitle": "测试论文",
        "author": "unknown",
        "totalChars": 100,
        "chunks": [
            {
                "chunkId": "chunk-0001",
                "sectionType": "body",
                "text": "综上所述，该研究具有重要意义，同时为后续研究提供参考。",
                "charCount": 28,
            }
        ],
    }
    results = [
        ai_check.Result(
            chunkId="chunk-0001",
            sensitivity="medium",
            aiScore=0.612,
            verdict="suspicious",
            reason="Generic closing claim.",
            signals=["template phrase", "generic claims"],
        )
    ]

    rows = ai_check.get_report_rows(scan, results)
    summary = ai_check.calculate_summary(scan, rows)
    rendered = ai_check.render_html(scan, rows, summary)

    assert "AI score: 0.612" in rendered
    assert "<th style=\"width:76%\">Text</th>" in rendered
    assert "Reason: Generic closing claim. Signals: template phrase, generic claims." in rendered
    assert "easyidea" not in rendered.lower()
    assert summary["aiScore"] > 0 and "aiRate" not in summary


def test_generate_report_writes_polished_pdf_without_libreoffice(tmp_path):
    scan = {
        "documentTitle": "测试论文",
        "author": "unknown",
        "totalChars": 100,
        "chunks": [
            {
                "chunkId": "chunk-0001",
                "sectionType": "body",
                "text": "综上所述，该研究具有重要意义，同时为后续研究提供参考。",
                "charCount": 28,
            }
        ],
    }
    results = {
        "results": [
            {
                "chunkId": "chunk-0001",
                "sensitivity": "medium",
                "aiScore": 0.612,
                "verdict": "suspicious",
                "reason": "泛化收束明显",
                "signals": ["模板表达", "泛化套话"],
            }
        ]
    }
    scan_path = tmp_path / "scan.json"
    results_path = tmp_path / "results.json"
    out_dir = tmp_path / "report"
    scan_path.write_text(json.dumps(scan, ensure_ascii=False), encoding="utf-8")
    results_path.write_text(json.dumps(results, ensure_ascii=False), encoding="utf-8")

    output = ai_check.generate_report(scan_path, results_path, out_dir)

    assert output["pdfError"] == ""
    assert Path(output["pdfPath"]).exists()
    assert Path(output["htmlPath"]).exists()
    assert Path(output["summaryPath"]).exists()
    assert Path(output["pdfPath"]).stat().st_size > 1000


def test_rewrite_validation_preserves_citations():
    item = ai_rewrite.Replacement(sourceText="已有方法在实验中取得提升[4]。", replacementText="已有方法在实验中取得提升。")
    assert ai_rewrite.validate_replacement_citations(item) == ["[4]"]

    ok = ai_rewrite.Replacement(
        sourceText="已有方法在实验中取得提升[4]。",
        replacementText="在相同实验条件下，已有方法取得了提升[4]。",
    )
    assert ai_rewrite.validate_replacement_citations(ok) == []


# ---------------------------------------------------------------------------
# scanner regressions
# ---------------------------------------------------------------------------


def test_paragraphs_starting_with_numbers_are_not_dropped():
    chunks = ai_check.build_chunks(
        "\n".join(
            [
                "1 绪论",
                "2019年以来，深度学习方法在水声目标识别任务上取得了显著进展，但模型在低信噪比条件下仍然难以稳定泛化。",
                "表2给出了不同方法在公开数据集上的识别准确率，可以看出本文方法在低信噪比条件下优势更加明显。",
                "表3 实验参数设置",
            ]
        )
    )
    text = "\n".join(chunk.text for chunk in chunks)
    assert "2019年以来" in text
    assert "表2给出了" in text
    assert "表3 实验参数设置" not in text
    assert all(chunk.sectionTitle == "1 绪论" for chunk in chunks)


def test_english_text_is_split_and_signals_are_found():
    features = ai_check.scan_trace_features(
        "In today's rapidly evolving landscape, deep learning plays a pivotal role in sonar analysis. "
        "Moreover, it is worth noting that the proposed framework is robust. "
        "Furthermore, the results underscore its significant potential. "
        "In conclusion, this work paves the way for future research."
    )
    assert features.sentenceCount == 4
    assert any(signal.startswith("template phrase") for signal in features.riskSignals)
    assert any(signal.startswith("generic claims") for signal in features.riskSignals)
    assert features.connectorCount >= 2


def test_report_fragments_keep_spaces_between_english_sentences():
    fragments = ai_check.build_report_fragments("First sentence is here. Second sentence follows it. Third one ends.")
    assert fragments == ["First sentence is here. Second sentence follows it. Third one ends."]


def test_results_accept_percentages_and_uppercase_enums():
    results = ai_check.normalize_results([{"chunkId": "chunk-0001", "aiScore": 85, "verdict": "AI", "sensitivity": "Medium"}])
    assert results[0].aiScore == 0.85
    assert results[0].verdict == "ai"


# ---------------------------------------------------------------------------
# DOCX round trips
# ---------------------------------------------------------------------------

from docx import Document  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls, qn  # noqa: E402

SOURCE = "综上所述，该方法具有重要意义，准确率提升了 2.1%[4]。"
REPLACEMENT = "在本文设定的低信噪比场景下，该方法的准确率提升了 2.1%[4]。"


def make_rich_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("摘要")
    paragraph = document.add_paragraph()
    paragraph.add_run("开头保留。").bold = True
    paragraph.add_run("综上所述，该方法具有重要意义，准确率提升了 2.1%")
    citation = paragraph.add_run("[4]")
    citation.font.superscript = True
    paragraph.add_run("。")
    paragraph._p.append(parse_xml(f'<m:oMath {nsdecls("m")}><m:r><m:t>x</m:t></m:r></m:oMath>'))
    paragraph.add_run("公式之后的文字。")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格里的文字：首先其次最后，这段文字不应进入默认扫描。"
    document.save(str(path))


def write_replacements(path: Path, items: list[dict]) -> Path:
    path.write_text(json.dumps({"replacements": items}, ensure_ascii=False), encoding="utf-8")
    return path


def body_paragraph(path: Path, index: int = 1):
    return Document(str(path)).paragraphs[index]._p


def test_scan_reads_document_order_and_skips_tables(tmp_path):
    source = tmp_path / "rich.docx"
    make_rich_docx(source)
    text, _meta = ai_check.read_docx_text(source)
    assert "表格里的文字" not in text
    assert text.index("开头保留") < text.index("公式之后的文字")
    with_tables, _meta = ai_check.read_docx_text(source, include_tables=True)
    assert "表格里的文字" in with_tables


def test_apply_preserves_formatting_math_and_citation(tmp_path):
    source = tmp_path / "rich.docx"
    make_rich_docx(source)
    out = tmp_path / "out.docx"
    result = ai_rewrite.apply_replacements(
        source, write_replacements(tmp_path / "r.json", [{"chunkId": "c1", "sourceText": SOURCE, "replacementText": REPLACEMENT}]), out
    )
    assert result["appliedCount"] == 1, result
    p = body_paragraph(out)
    text = ai_check.paragraph_element_text(p)
    assert text == "开头保留。" + REPLACEMENT + "公式之后的文字。"
    assert p.find(qn("m:oMath")) is not None, "inline equation must survive"
    runs = p.findall(qn("w:r"))
    bold = [r for r in runs if r.find(f"{qn('w:rPr')}/{qn('w:b')}") is not None]
    assert "".join(t.text for r in bold for t in r.findall(qn("w:t"))) == "开头保留。"
    sup = [r for r in runs if r.find(f"{qn('w:rPr')}/{qn('w:vertAlign')}") is not None]
    assert "".join(t.text for r in sup for t in r.findall(qn("w:t"))) == "[4]"
    assert len(runs) < 12, "single-character runs should be merged back"
    assert source.read_bytes() != out.read_bytes()


def test_apply_with_track_changes_keeps_original_recoverable(tmp_path):
    source = tmp_path / "rich.docx"
    make_rich_docx(source)
    out = tmp_path / "tracked.docx"
    result = ai_rewrite.apply_replacements(
        source,
        write_replacements(tmp_path / "r.json", [{"sourceText": SOURCE, "replacementText": REPLACEMENT}]),
        out,
        track_changes=True,
    )
    assert result["appliedCount"] == 1 and result["trackChanges"] is True
    p = body_paragraph(out)
    assert p.findall(qn("w:ins")) and p.findall(qn("w:del"))
    accepted = ai_check.paragraph_element_text(p)
    assert accepted == "开头保留。" + REPLACEMENT + "公式之后的文字。"
    # rejecting = drop w:ins content, keep w:delText: must give back the original sentence
    rejected = ""
    for element in p.iter(qn("w:t"), qn("w:delText")):
        inside_ins = any(a.tag == qn("w:ins") for a in element.iterancestors())
        if not inside_ins:
            rejected += element.text or ""
    assert SOURCE in rejected
    for wrapper in p.iter(qn("w:ins"), qn("w:del")):
        assert wrapper.get(qn("w:author")) == "AI-check" and wrapper.get(qn("w:id"))


def test_apply_refuses_to_overwrite_source(tmp_path):
    source = tmp_path / "rich.docx"
    make_rich_docx(source)
    replacements = write_replacements(tmp_path / "r.json", [{"sourceText": SOURCE, "replacementText": REPLACEMENT}])
    try:
        ai_rewrite.apply_replacements(source, replacements, source)
    except ValueError as exc:
        assert "never modifies the original" in str(exc)
    else:
        raise AssertionError("expected ValueError")


def test_apply_skips_rewrites_that_drop_numbers(tmp_path):
    source = tmp_path / "rich.docx"
    make_rich_docx(source)
    out = tmp_path / "out.docx"
    replacements = write_replacements(
        tmp_path / "r.json", [{"sourceText": SOURCE, "replacementText": "该方法在实验中表现更好，准确率有所提升[4]。"}]
    )
    result = ai_rewrite.apply_replacements(source, replacements, out)
    assert result["appliedCount"] == 0
    assert "missing numbers: 2.1" in result["skipped"][0]["reason"]


def test_scan_reads_latex_project_directory(tmp_path):
    (tmp_path / "sections").mkdir()
    (tmp_path / "main.tex").write_text(
        "\\documentclass{article}\n"
        "\\title{A \\emph{Test} Paper}\n"
        "\\newcommand{\\tsection}[2]{\\section*{#1}\\label{#2}}\n"
        "\\begin{document}\n"
        "\\noindent\\textbf{Abstract -}\n"
        "We study $O(n^2)$ algorithms for nested {braces} and report a 50\\% speedup~\\citep[p.~3]{a,b}. % hidden-comment\n"
        "Results hold for $\\alpha > 0$ on every benchmark we tried in this study.\n\n"
        "\\noindent\\textbf{Keywords --} verification, LaTeX.\n\n"
        "\\input{sections/intro}\n"
        "\\bibliography{refs}\n"
        "\\end{document}\n",
        encoding="utf-8",
    )
    (tmp_path / "sections" / "intro.tex").write_text(
        "\\tsection{Introduction}{sec:intro}\n"
        "\\paragraph{Motivation.} Engineers choose verification scopes by hand, as noted by \\citet{c}.\\todo{cite more}\n"
        "\\begin{figure}\\caption{Hidden caption text}\\end{figure}\n"
        "\\begin{tabular}{ll} Cell one text here & cell two text here \\\\ row two & more \\end{tabular}\n",
        encoding="utf-8",
    )

    data = ai_check.scan_source(tmp_path)
    chunks = data["chunks"]
    texts = " ".join(chunk["text"] for chunk in chunks)

    assert data["documentTitle"] == "A Test Paper"
    assert [chunk["sectionType"] for chunk in chunks] == ["abstract", "body"]
    assert chunks[1]["sectionTitle"] == "1.1 Motivation"
    assert "speedup [1, 2]" in texts and "Results hold for alpha > 0" in texts
    assert "as noted by [3]" in texts
    for dropped in ("hidden-comment", "Keywords", "cite more", "Hidden caption", "Cell one", "refs"):
        assert dropped not in texts


def test_scan_latex_include_tables_keeps_rows_as_paragraphs(tmp_path):
    source = tmp_path / "body.tex"
    source.write_text(
        "\\section{Results}\n"
        "Our evaluation covers two open-source accelerator designs in full detail.\n\n"
        "\\begin{tabular}{ll} Month one establishes the baseline configurations & done \\\\ "
        "Month two integrates the contract generation loop & planned \\end{tabular}\n",
        encoding="utf-8",
    )

    without_tables = " ".join(chunk["text"] for chunk in ai_check.scan_source(source)["chunks"])
    with_tables = " ".join(chunk["text"] for chunk in ai_check.scan_source(source, include_tables=True)["chunks"])

    assert "Month one" not in without_tables
    assert "Month one establishes the baseline configurations done" in with_tables
    assert "Our evaluation covers two open-source accelerator designs" in with_tables


def test_scan_missing_latex_source_reports_path(tmp_path):
    missing = tmp_path / "missing.tex"
    try:
        ai_check.scan_source(missing)
    except RuntimeError as exc:
        assert "source not found" in str(exc)
    else:
        raise AssertionError("expected RuntimeError")


def test_results_accept_legacy_aigc_value_key():
    results = ai_check.normalize_results([{"chunkId": "chunk-0001", "aigcValue": 0.4}])
    assert results[0].aiScore == 0.4


def test_scan_reads_markdown_headings_lists_and_skips_code(tmp_path):
    source = tmp_path / "notes.md"
    source.write_text(
        "---\ntitle: ignored front matter\n---\n"
        "## Background\n\n"
        "Engineers choose **verification scopes** by hand, and the [manual](http://x.y) effort grows with each design.\n"
        "The same paragraph continues on a soft-wrapped line here.\n\n"
        "```python\nprint('hidden code block')\n```\n\n"
        "| Col | Other |\n|-----|-------|\n| hidden table cell text | x |\n\n"
        "- A list item that is long enough to survive the minimum fragment length filter.\n",
        encoding="utf-8",
    )

    chunks = ai_check.scan_source(source)["chunks"]
    texts = " ".join(chunk["text"] for chunk in chunks)

    assert chunks[0]["sectionTitle"] == "1 Background"
    assert "Engineers choose verification scopes by hand, and the manual effort" in texts
    assert "soft-wrapped line here." in texts
    assert "A list item that is long enough" in texts
    for dropped in ("front matter", "hidden code", "hidden table", "**", "http://"):
        assert dropped not in texts


def test_scan_reads_plain_text_and_raw_text_as_prose():
    text = (
        "This email explains why the deployment slipped by two days last week.\n"
        "The root cause was a misconfigured staging cluster.\n\n"
        "We will add a pre-flight check so the same mistake cannot reach production again."
    )
    data = ai_check.scan_source(None, text=text, title="Status email")
    texts = " ".join(chunk["text"] for chunk in data["chunks"])

    assert data["documentTitle"] == "Status email"
    assert "deployment slipped by two days last week. The root cause" in texts
    assert "pre-flight check" in texts


def test_scan_reads_html_without_scripts(tmp_path):
    source = tmp_path / "post.html"
    source.write_text(
        "<html><head><title>x</title><script>var hidden = 1;</script></head><body>"
        "<h2>Results</h2><p>Our method reduced verification time on both accelerator designs we tested.</p>"
        "<p>It still fails on designs with multiple clock domains, which we leave to future work.</p>"
        "</body></html>",
        encoding="utf-8",
    )
    chunks = ai_check.scan_source(source)["chunks"]
    texts = " ".join(chunk["text"] for chunk in chunks)

    assert chunks[0]["sectionTitle"] == "1 Results"
    assert "reduced verification time" in texts and "multiple clock domains" in texts
    assert "hidden" not in texts


def test_scan_reads_pdf_text(tmp_path):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    source = tmp_path / "paper.pdf"
    pdf = canvas.Canvas(str(source), pagesize=A4)
    pdf.drawString(72, 760, "The prototype checks each module against its contract before composing proofs.")
    pdf.save()

    texts = " ".join(chunk["text"] for chunk in ai_check.scan_source(source)["chunks"])
    assert "checks each module against its contract" in texts
